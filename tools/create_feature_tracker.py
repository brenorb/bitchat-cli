from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "bitchat-cli-feature-tracker.xlsx"


ROWS = [
    ("CLI-001", "Core", "doctor", "As a user, I can inspect local configuration and scope.", "Prints configuration and supported scope; --json is machine-readable.", "Not tested", "tests/test_app.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-002", "Chat", "chat", "As a user, I can start an interactive BLE public chat.", "Starts transport, shows messages, accepts input, and shuts down cleanly.", "Not tested", "tests/test_app.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-003", "Chat", "chat public message", "As a chat user, I can send non-slash text publicly.", "Text becomes a signed public BitChat message and is broadcast to connected peers.", "Not tested", "tests/test_app.py; tests/test_transport.py", "README.md"),
    ("CLI-004", "Chat", "/help", "As a chat user, I can discover available chat commands.", "Displays the command help without sending a network message.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-005", "Chat", "/name", "As a chat user, I can change my nickname.", "Updates local identity display and broadcasts an announce.", "Not tested", "tests/test_transport.py; src/bitchat_cli/transport.py", "README.md"),
    ("CLI-006", "Chat", "/who", "As a chat user, I can see known peers.", "Lists current peer nicknames and identifiers.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-007", "Chat", "/msg", "As a chat user, I can send a direct message.", "Resolves nickname or peer ID, performs Noise handshake if needed, and delivers the DM.", "Not tested", "tests/test_transport.py; src/bitchat_cli/transport.py", "README.md"),
    ("CLI-008", "Chat", "/clear", "As a chat user, I can clear the visible chat.", "Clears the local interactive display only.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-009", "Chat", "/quit", "As a chat user, I can leave chat safely.", "Stops BLE, sends leave, and releases the radio lease.", "Not tested", "tests/test_transport.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-010", "Daemon", "daemon start", "As a user, I can run a background listener.", "Starts one daemon, persists its PID/socket, and records inbound traffic.", "Not tested", "tests/test_app.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-011", "Daemon", "daemon status", "As a user, I can inspect daemon health.", "Shows process, links, peers, advertising, subscribers, relays, and watches.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-012", "Daemon", "daemon stop", "As a user, I can stop the daemon cleanly.", "Sends leave, waits for exit, and cleans PID/radio runtime state.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-013", "Inbox", "inbox", "As a user, I can read recorded public messages.", "Shows persisted public messages with --tail and --json options.", "Not tested", "src/bitchat_cli/queries.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-014", "Mesh", "mesh send", "As a user, I can explicitly send to the #mesh conversation.", "Routes a public message through the daemon's BLE transport.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-015", "Mesh", "mesh inbox", "As a user, I can read the #mesh conversation.", "Shows the same BLE public history as inbox.", "Not tested", "src/bitchat_cli/queries.py", "README.md"),
    ("CLI-016", "Peers", "peers", "As a user, I can inspect learned peers.", "Shows signed peer records, addresses, timestamps, and optional JSON.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/queries.py", "README.md"),
    ("CLI-017", "Files", "send-file", "As a user, I can broadcast a file.", "Encodes the official file TLV, fragments if needed, and persists receivers' files.", "Not tested", "tests/test_transport.py; src/bitchat_cli/files.py", "README.md"),
    ("CLI-018", "Files", "mesh send-file", "As a user, I can send a file explicitly to #mesh.", "Uses the same public file path as send-file through the mesh daemon.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/files.py", "README.md"),
    ("CLI-019", "DM", "dm send", "As a user, I can send a direct message through the daemon.", "Uses BLE first and queues/falls back only according to social policy.", "Not tested", "tests/test_app.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-020", "DM", "dm send-file", "As a user, I can send a file privately.", "Uses Noise encryption and persists the incoming attachment with metadata.", "Not tested", "tests/test_transport.py; src/bitchat_cli/files.py", "README.md"),
    ("CLI-021", "DM", "dm inbox", "As a user, I can list DM conversations.", "Groups and summarizes recorded private conversations.", "Not tested", "src/bitchat_cli/queries.py", "README.md"),
    ("CLI-022", "DM", "dm history", "As a user, I can inspect one DM history.", "Resolves a peer and prints its persisted messages with tail/JSON options.", "Not tested", "src/bitchat_cli/queries.py", "README.md"),
    ("CLI-023", "Geo", "geo send", "As a user, I can publish to a geohash channel.", "Publishes a signed kind 20000 event with optional PoW, teleport, and privacy acknowledgement.", "Not tested", "tests/test_geochat.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-024", "Geo", "geo inbox", "As a user, I can read one geohash channel.", "Fetches and renders deduplicated events with PoW/teleport markers.", "Not tested", "tests/test_geochat.py; src/bitchat_cli/queries.py", "README.md"),
    ("CLI-025", "Geo", "geo who", "As a user, I can see geohash participants.", "Lists recent participants from messages and presence, excluding self.", "Not tested", "src/bitchat_cli/queries.py", "README.md"),
    ("CLI-026", "Geo", "geo levels", "As a user, I can derive the hierarchy of a geohash.", "Prints region through building prefixes offline.", "Passed", "tests/test_geochat.py; src/bitchat_cli/geochat.py", "README.md"),
    ("CLI-027", "Geo", "geo encode", "As a user, I can encode coordinates as geohashes.", "Prints a selected level or all supported levels offline.", "Passed", "tests/test_geochat.py; src/bitchat_cli/geochat.py", "README.md"),
    ("CLI-028", "Geo", "geo decode", "As a user, I can decode a geohash.", "Prints center and bounds offline.", "Passed", "tests/test_geochat.py; src/bitchat_cli/geochat.py", "README.md"),
    ("CLI-029", "Geo", "geo lookup", "As a user, I can resolve a place name.", "Queries a third-party geocoder without persisting the query or coordinates.", "Not tested", "src/bitchat_cli/geochat.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-030", "Geo metadata", "geo channel set", "As a user, I can save local channel metadata.", "Stores name, note, bookmark, selected context, and teleport flags.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/state.py", "README.md"),
    ("CLI-031", "Geo metadata", "geo channel show", "As a user, I can inspect saved channel metadata.", "Shows metadata and recent activity for one geohash.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-032", "Geo metadata", "geo channel list", "As a user, I can list saved channels.", "Lists all or only bookmarked local geohash metadata.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-033", "Location notes", "geo notes add", "As a user, I can publish a building location note.", "Publishes a persistent kind 1 note at building precision.", "Not tested", "tests/test_geochat.py; src/bitchat_cli/location_notes.py", "README.md"),
    ("CLI-034", "Location notes", "geo notes list", "As a user, I can fetch building location notes.", "Fetches and limits persistent notes for one building geohash.", "Not tested", "src/bitchat_cli/location_notes.py", "README.md"),
    ("CLI-035", "Geo trust", "geo participant show", "As a user, I can inspect geohash participant trust.", "Shows saved favorite, blocked, and verified flags.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/queries.py", "README.md"),
    ("CLI-036", "Geo trust", "geo participant block/unblock", "As a user, I can moderate a geohash participant.", "Persists block state and removes it on unblock.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-037", "Geo DM", "geo dm send", "As a user, I can send a geohash-scoped DM.", "Uses the deterministic per-geohash identity and embedded BitChat payload.", "Not tested", "tests/test_geochat.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-038", "Geo DM", "geo dm history", "As a user, I can inspect a geohash-scoped DM history.", "Shows persisted history for the selected participant.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/queries.py", "README.md"),
    ("CLI-039", "Geo watch", "geo watch add/remove/list", "As a user, I can manage background geohash subscriptions.", "Maintains relay subscriptions and periodic coarse presence for watched channels.", "Not tested", "tests/test_geochat.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-040", "Social", "social list", "As a user, I can inspect social mappings.", "Lists linked Nostr identities and trust flags.", "Not tested", "src/bitchat_cli/social.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-041", "Social", "social link/unlink", "As a user, I can link or remove a peer's Nostr identity.", "Persists or removes the mapping by peer ID or nickname.", "Not tested", "src/bitchat_cli/social.py", "README.md"),
    ("CLI-042", "Social", "social favorite/unfavorite", "As a user, I can authorize or disable Nostr fallback.", "Toggles the favorite flag without changing the link itself.", "Not tested", "src/bitchat_cli/social.py", "README.md"),
    ("CLI-043", "Social", "social block/unblock", "As a user, I can block or unblock a linked peer.", "Toggles fallback blocking and prevents disallowed remote delivery.", "Not tested", "src/bitchat_cli/social.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-044", "Social", "social fingerprint", "As a user, I can inspect peer trust material.", "Shows stable cryptographic fingerprints for a resolved peer.", "Not tested", "src/bitchat_cli/verification.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-045", "Social", "social verify/unverify", "As a user, I can mark a peer verified or clear verification.", "Validates an optional verification payload before persisting verified state.", "Not tested", "tests/test_noise.py; src/bitchat_cli/verification.py", "README.md"),
    ("CLI-046", "Network", "network proxy status", "As a user, I can inspect privacy routing policy.", "Shows proxy mode, policy, and resolved target.", "Not tested", "tests/test_nostr.py; src/bitchat_cli/network.py", "README.md"),
    ("CLI-047", "Network", "network proxy set/disable", "As a user, I can configure or disable outbound proxying.", "Persists tor, system, or explicit proxy URL with prefer/require policy.", "Not tested", "tests/test_nostr.py; src/bitchat_cli/network.py", "README.md"),
    ("CLI-048", "Network", "network proxy test", "As a user, I can test local proxy reachability.", "Reports whether the configured proxy is usable.", "Not tested", "src/bitchat_cli/network.py; src/bitchat_cli/app.py", "README.md"),
    ("CLI-049", "Topics", "topic join/leave", "As a user, I can manage a named local topic.", "Normalizes the channel, stores membership, and removes it on leave.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/state.py", "README.md"),
    ("CLI-050", "Topics", "topic list", "As a user, I can list configured topics.", "Shows local topic membership and metadata, optionally as JSON.", "Not tested", "src/bitchat_cli/app.py", "README.md"),
    ("CLI-051", "Topics", "topic password/retain", "As a user, I can configure topic password and retention markers.", "Stores, clears, enables, or disables local topic metadata.", "Not tested", "src/bitchat_cli/app.py; src/bitchat_cli/state.py", "README.md"),
    ("CLI-052", "Diagnostics", "selftest", "As a user, I can run offline protocol checks.", "Runs crypto, packet, Noise, fragmentation, and state checks without BLE or network.", "Passed", "tests/; src/bitchat_cli/app.py", "README.md"),
    ("CLI-053", "BLE presence", "daemon idle presence", "As a connected mobile user, I continue seeing clanker while the CLI is idle.", "CLI sends a signed announce every 30s while it has a live client or peripheral subscriber; Android still lists clanker after >60s idle.", "Passed", "20d5577; tests/test_transport.py; src/bitchat_cli/transport.py", "Android upstream: BluetoothMeshService.kt; iOS upstream: BLEService.swift"),
    ("CLI-054", "BLE interop", "Android ↔ CLI public chat", "As trueanon, I can receive a message from clanker; as clanker, I can receive a message from trueanon.", "Both directions render the sender nickname and message exactly once.", "Passed", "tests/test_app.py; tests/test_transport.py; manual 2026-07-25", "Android package com.bitchat.droid; Android_to_cli_1900; cli_to_android_1921"),
    ("CLI-055", "BLE interop", "Android ↔ CLI DM", "As either peer, I can send and receive a private message.", "Noise handshake completes, ciphertext is delivered, and DM history records it once.", "Passed", "tests/test_noise.py; tests/test_transport.py; manual 2026-07-25", "dm_cli_to_android_1926; Android_to_cli_dm_1927; Android package com.bitchat.droid"),
    ("CLI-056", "BLE interop", "Android ↔ CLI file", "As either peer, I can send a file over BLE.", "Large payloads fragment/reassemble and persist with correct name, size, and MIME type.", "Passed", "tests/test_fragments.py; tests/test_transport.py; manual 2026-07-25", "bitchat-test-file.txt; Android photo img_1785004221247.jpg; Android package com.bitchat.droid"),
    ("CLI-057", "Sync", "reconnect sync", "As a reconnecting peer, I receive missed public packets without duplicates.", "REQUEST_SYNC/RSR exchanges recover supported packets and persistence stays deduplicated.", "Passed", "tests/test_sync.py; tests/test_transport.py; manual 2026-07-25", "Sync_gap_android_1934 recovered after daemon restart; messages.jsonl count=1"),
    ("CLI-058", "Nostr fallback", "offline DM fallback", "As a linked and favorited peer, I can receive an offline text DM through Nostr.", "BLE is preferred; fallback is used only when policy and trust state allow it.", "Blocked", "tests/test_nostr.py; src/bitchat_cli/app.py; manual 2026-07-25", "CLI routed nostr_fallback_cli_1938 via Nostr; Android UI did not surface it during bounded observation"),
]


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Features"
    headers = ["ID", "Area", "Command / flow", "User story", "Expected behavior", "Manual status", "Automated/code evidence", "Reference"]
    ws.append(headers)
    for row in ROWS:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12, 18, 28, 52, 86, 20, 46, 48]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    validation = DataValidation(type="list", formula1='"Not tested,Passed,Failed,Blocked,Not applicable"', allow_blank=False)
    ws.add_data_validation(validation)
    validation.add(f"F2:F{ws.max_row}")
    ws.conditional_formatting.add(f"F2:F{ws.max_row}", FormulaRule(formula=['F2="Passed"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"F2:F{ws.max_row}", FormulaRule(formula=['F2="Failed"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(f"F2:F{ws.max_row}", FormulaRule(formula=['F2="Blocked"'], fill=PatternFill("solid", fgColor="FFEB9C")))

    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value", "Definition"])
    summary.append(["Total user stories", "=COUNTA(Features!A2:A59)", "All tracked CLI features and interoperability flows"])
    summary.append(["Passed", "=COUNTIF(Features!F2:F59,\"Passed\")", "Manual or offline behavior verified"])
    summary.append(["Failed", "=COUNTIF(Features!F2:F59,\"Failed\")", "Observed behavior contradicts the story"])
    summary.append(["Blocked", "=COUNTIF(Features!F2:F59,\"Blocked\")", "Could not test because of an external blocker"])
    summary.append(["Not tested", "=COUNTIF(Features!F2:F59,\"Not tested\")", "Still needs a test run"])
    summary.append(["Coverage", "=IFERROR(B3/B2,0)", "Passed stories divided by total stories"])
    summary.append([])
    summary.append(["Canonical references", "URL / path", "Purpose"])
    summary.append(["Android upstream", "https://github.com/permissionlesstech/bitchat-android", "Reference implementation and BLE behavior"])
    summary.append(["iOS/macOS upstream", "https://github.com/permissionlesstech/bitchat", "Reference implementation and BLE behavior"])
    summary.append(["Current fix commit", "20d5577", "Periodic BLE presence announcements"])
    summary.freeze_panes = "A2"
    for cell in summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 64
    summary.column_dimensions["C"].width = 64
    summary["B7"].number_format = "0.0%"

    errors = wb.create_sheet("Errors")
    errors.append(["Error ID", "Date", "Area", "Observed error", "Impact", "Status", "Evidence", "Next action"])
    errors.append(["ERR-001", "2026-07-25", "Nostr relays", "Several configured relays returned HTTP 503 or timed out during daemon startup.", "Geo relay connectivity is degraded; BLE presence test remained independent.", "External dependency", "daemon.log during manual test", "Retest relay health separately; do not conflate with BLE failures."])
    errors.append(["ERR-002", "2026-07-25", "BLE presence", "clanker disappeared from Android while CLI was idle.", "Peer was not visible in the mobile room after the peer liveness window.", "Fixed", "20d5577; Android UI showed clanker after >60s idle", "Run full post-fix interoperability matrix."])
    errors.append(["ERR-003", "2026-07-25", "Test procedure", "The first inbox query ran before the asynchronous Android write had been persisted.", "The message was temporarily absent from the CLI query even though it arrived moments later.", "Test procedure", "messages.jsonl and daemon.log show Android_to_cli_1900 at 18:19:57/18:21:00", "Wait for delivery or poll with a bounded timeout before declaring a delivery failure."])
    errors.append(["ERR-004", "2026-07-25", "Nostr fallback", "The CLI returned via nostr [nostr_fallback] and persisted nostr_fallback_cli_1938, but the Android UI did not show the DM during the bounded observation window.", "Offline cross-platform DM delivery is not confirmed; BLE behavior remains unaffected.", "Open", "messages.jsonl message_id nostr-1785004686603; Android UI/logcat; relay availability warnings", "Compare Android Nostr subscription/identity handling with upstream and repeat with a known-current npub."])
    errors.freeze_panes = "A2"
    errors.auto_filter.ref = errors.dimensions
    for index, width in enumerate([14, 14, 18, 72, 58, 22, 44, 56], 1):
        errors.column_dimensions[chr(64 + index)].width = width
    for cell in errors[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in errors.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
